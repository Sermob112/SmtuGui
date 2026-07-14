import csv
import sqlite3
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill

SOURCES_DIR = "sources"
OUTPUT_XLSX = "result.xlsx"
DB_PATH = r"C:\Users\Sergey\Desktop\Work\SmtuGui\smtuIdle\database.db"

KEY_COLUMN_NAME = "Реестровый номер закупки"
KEY_COLUMN_INDEX = 1  # второй столбец (индексация с 0)

PRICE_COLUMN_INDEX = 8      # столбец I - Начальная максимальная цена
CURRENCY_COLUMN_INDEX = 9   # столбец J - Валюта
OKPD2_COLUMN_INDEX = 14     # столбец O - Классификация по ОКПД2

MIN_PRICE = 20_000_000
TARGET_CURRENCY = "RUB"
ALLOWED_OKPD2_PREFIXES = ( "30", "77", "64", "52", "42", "41")

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

skipped_files = []
processed_files = []
seen_keys = set()
header_written = None
rows_to_write = []
def normalize_registry_number(value):
    """Убирает символ '№' и лишние пробелы для корректного сравнения с БД."""
    if not value:
        return ""
    return value.replace('№', '').strip()

def parse_price(value):
    """Преобразует строку цены в float, учитывая пробелы как разделители тысяч и запятую как десятичный разделитель."""
    if value is None:
        return None
    value = value.strip().replace('\xa0', '').replace(' ', '').replace(',', '.')
    if not value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def okpd2_matches(value):
    """Проверяет, начинается ли код ОКПД2 с одного из разрешённых двузначных префиксов.
    Код может быть вида '72.19.50.000' или '30.11.33.190:Суда прочии'."""
    if not value:
        return False
    code = value.strip().split(':', 1)[0]
    code_prefix = code.split('.', 1)[0]
    return code_prefix in ALLOWED_OKPD2_PREFIXES


def load_registry_numbers_from_db(db_path):
    """Читает все значения RegistryNumber из таблицы purchase (peewee создаёт имя таблицы в нижнем регистре)."""
    registry_numbers = set()
    conn = sqlite3.connect(db_path)
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")
        tables = [row[0] for row in cursor.fetchall()]

        table_name = None
        for t in tables:
            if t.lower() == "purchase":
                table_name = t
                break

        if table_name is None:
            print(f"Таблица purchase не найдена. Доступные таблицы: {tables}")
            return registry_numbers

        cursor.execute(f'SELECT RegistryNumber FROM "{table_name}"')
        for (value,) in cursor.fetchall():
            if value:
                registry_numbers.add(normalize_registry_number(value))
    finally:
        conn.close()
    return registry_numbers


source_path = Path(SOURCES_DIR)
csv_files = sorted(source_path.glob("*.csv"))

print(f"Найдено файлов: {len(csv_files)}")
def detect_delimiter(file_path, encoding='windows-1251'):
    with open(file_path, 'r', encoding=encoding, newline='') as f:
        sample = f.read(2048)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=';,')
        return dialect.delimiter
    except csv.Error:
        return ';'  # запасной вариант по умолчанию
for file_path in csv_files:
    try:
        delim = detect_delimiter(file_path)
        with open(file_path, 'r', encoding='windows-1251', newline='') as f:
            reader = csv.reader(f, delimiter=delim)
            try:
                header = next(reader)
            except StopIteration:
                skipped_files.append((file_path.name, "пустой файл"))
                continue

            if len(header) <= KEY_COLUMN_INDEX:
                skipped_files.append((file_path.name, "нет второго столбца"))
                continue

            if header[KEY_COLUMN_INDEX].strip() != KEY_COLUMN_NAME:
                skipped_files.append(
                    (file_path.name, f"неверное имя столбца: '{header[KEY_COLUMN_INDEX]}'")
                )
                continue

            required_max_index = max(PRICE_COLUMN_INDEX, CURRENCY_COLUMN_INDEX, OKPD2_COLUMN_INDEX)
            if len(header) <= required_max_index:
                skipped_files.append((file_path.name, "нет столбца цены/валюты/ОКПД2"))
                continue

            if header_written is None:
                header_written = header

            if len(header) != len(header_written):
                skipped_files.append((file_path.name, "иное количество столбцов"))
                continue

            file_row_count = 0
            for row in reader:
                if len(row) <= required_max_index:
                    continue

                key = row[KEY_COLUMN_INDEX].strip()
                if not key or key in seen_keys:
                    continue

                if not okpd2_matches(row[OKPD2_COLUMN_INDEX]):
                    continue

                currency = row[CURRENCY_COLUMN_INDEX].strip().upper()

                if currency == TARGET_CURRENCY:
                    price = parse_price(row[PRICE_COLUMN_INDEX])
                    if price is None or price < MIN_PRICE:
                        continue

                seen_keys.add(key)
                rows_to_write.append(row)
                file_row_count += 1

            processed_files.append((file_path.name, file_row_count))

    except UnicodeDecodeError:
        skipped_files.append((file_path.name, "ошибка декодирования windows-1251"))
    except Exception as e:
        skipped_files.append((file_path.name, f"ошибка: {e}"))

# Загружаем реестровые номера из БД
db_registry_numbers = load_registry_numbers_from_db(DB_PATH)
print(f"Загружено записей из БД: {len(db_registry_numbers)}")

# Запись результата в XLSX с подсветкой
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Result"

if header_written:
    ws.append(header_written)

matched_count = 0
for row in rows_to_write:
    ws.append(row)
    key_value = normalize_registry_number(row[KEY_COLUMN_INDEX])
    if key_value in db_registry_numbers:
        matched_count += 1
        excel_row = ws.max_row
        for col_idx in range(1, len(row) + 1):
            ws.cell(row=excel_row, column=col_idx).fill = GREEN_FILL

wb.save(OUTPUT_XLSX)

# Отчёт
print(f"\nОбработано файлов: {len(processed_files)}")
print(f"Пропущено файлов: {len(skipped_files)}")
print(f"Уникальных записей в итоге (после фильтров): {len(rows_to_write)}")
print(f"Совпало с базой данных (подсвечено зелёным): {matched_count}")

if skipped_files:
    print("\nПропущенные файлы и причины:")
    for name, reason in skipped_files:
        print(f"  {name}: {reason}")